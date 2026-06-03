"""Cases 22-24: chart types now available through local Witan Alfred.

The local Witan Alfred xlsx-serve binary can author chart surfaces that were
not available through the public npx witan CLI when the earlier probes were
run:

- waterfall charts with subtotal/total markers
- stock OHLC charts
- column charts with multi-level string category labels

This script compares Witan Alfred as the control against direct openpyxl
authoring and openpyxl followed by a LibreOffice save.
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
from pathlib import Path
from typing import Any

import openpyxl.chart as openpyxl_charts
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, StockChart


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "beta_chart_types"
SOFFICE = Path(os.environ.get("SOFFICE", "/opt/homebrew/bin/soffice"))
XLSX_SERVE = Path(
    os.environ.get(
        "XLSX_SERVE",
        str((ROOT / "../witan-alfred/bin/publish/xlsx-serve").resolve()),
    )
)


def run(cmd: list[str]) -> subprocess.CompletedProcess[str]:
    proc = subprocess.run(
        cmd,
        cwd=ROOT,
        check=False,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
    )
    if proc.returncode:
        raise RuntimeError(proc.stdout)
    return proc


def run_xlsx_serve(args: list[str]) -> Any:
    proc = run([str(XLSX_SERVE), *args, "--json"])
    payload = json.loads(proc.stdout)
    if not payload.get("ok"):
        raise RuntimeError(json.dumps(payload, indent=2))
    return payload.get("result")


def resave_with_libreoffice(src: Path) -> Path:
    if not SOFFICE.exists():
        raise RuntimeError(f"LibreOffice not found at {SOFFICE}; set SOFFICE=/path/to/soffice")

    work = OUT / "_lo_work"
    converted = OUT / "_lo_out"
    profile = OUT / "_lo_profile"
    shutil.rmtree(work, ignore_errors=True)
    shutil.rmtree(converted, ignore_errors=True)
    shutil.rmtree(profile, ignore_errors=True)
    work.mkdir(parents=True)
    converted.mkdir(parents=True)
    profile.mkdir(parents=True)

    local = work / src.name
    shutil.copy(src, local)
    run(
        [
            str(SOFFICE),
            "--headless",
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(converted),
            str(local),
        ]
    )
    out = converted / local.name
    if not out.exists():
        raise RuntimeError(f"LibreOffice did not produce {out}")
    return out


def get_chart(path: Path, sheet: str, name: str) -> dict[str, Any]:
    return run_xlsx_serve(
        [
            "exec",
            str(path),
            "--expr",
            f'xlsx.getChart(wb, "{sheet}", "{name}")',
        ]
    )


def pos_range(chart: dict[str, Any]) -> str:
    position = chart.get("position", {})
    start = position.get("from", {}).get("cell")
    end = position.get("to", {}).get("cell")
    return f"{start}:{end}"


def first_group(chart: dict[str, Any]) -> dict[str, Any]:
    groups = chart.get("groups", [])
    return groups[0] if groups else {}


def first_series(chart: dict[str, Any]) -> dict[str, Any]:
    return (first_group(chart).get("series") or [{}])[0]


def witan_waterfall() -> dict[str, Any]:
    dst = OUT / "case22_witan_waterfall.xlsx"
    code = r'''
await xlsx.addSheet(wb, "Summary");
await xlsx.setCells(wb, [
  {address:"Summary!A1", value:"Step"},       {address:"Summary!B1", value:"Amount"},
  {address:"Summary!A2", value:"Revenue"},    {address:"Summary!B2", value:1200},
  {address:"Summary!A3", value:"COGS"},       {address:"Summary!B3", value:-420},
  {address:"Summary!A4", value:"Payroll"},    {address:"Summary!B4", value:-250},
  {address:"Summary!A5", value:"Gross"},      {address:"Summary!B5", value:530},
  {address:"Summary!A6", value:"Tax"},        {address:"Summary!B6", value:-95},
  {address:"Summary!A7", value:"Net"},        {address:"Summary!B7", value:435}
]);
await xlsx.addChart(wb, "Summary", {
  name: "Profit Bridge",
  position: { from: { cell: "D2" }, to: { cell: "L18" } },
  groups: [{
    type: "waterfall",
    series: [{
      name: { ref: "Summary!B1" },
      categories: "Summary!A2:A7",
      values: "Summary!B2:B7",
      totalIndexes: [3, 5],
      dataLabels: { showValue: true, position: "outsideEnd" }
    }]
  }],
  title: { text: "Profit Bridge" },
  legend: { position: "topRight" },
  axes: { value: { numberFormat: "$#,##0" } },
  styleId: 395
});
return await xlsx.getChart(wb, "Summary", "Profit Bridge");
'''
    return run_xlsx_serve(["exec", str(dst), "--create", "--save", "--code", code])


def witan_stock() -> dict[str, Any]:
    dst = OUT / "case23_witan_stock.xlsx"
    code = r'''
await xlsx.addSheet(wb, "Prices");
await xlsx.setCells(wb, [
  {address:"Prices!A1", value:"Date"},       {address:"Prices!B1", value:"Open"}, {address:"Prices!C1", value:"High"}, {address:"Prices!D1", value:"Low"}, {address:"Prices!E1", value:"Close"},
  {address:"Prices!A2", value:"2026-01-01"}, {address:"Prices!B2", value:100},    {address:"Prices!C2", value:112},    {address:"Prices!D2", value:96},     {address:"Prices!E2", value:108},
  {address:"Prices!A3", value:"2026-01-02"}, {address:"Prices!B3", value:108},    {address:"Prices!C3", value:116},    {address:"Prices!D3", value:101},    {address:"Prices!E3", value:103},
  {address:"Prices!A4", value:"2026-01-03"}, {address:"Prices!B4", value:103},    {address:"Prices!C4", value:120},    {address:"Prices!D4", value:99},     {address:"Prices!E4", value:117}
]);
await xlsx.addChart(wb, "Prices", {
  name: "OHLC",
  position: { from: { cell: "G2" }, to: { cell: "N18" } },
  groups: [{
    type: "stockOHLC",
    series: [
      { name: { ref: "Prices!B1" }, stockRole: "open",  categories: "Prices!A2:A4", values: "Prices!B2:B4" },
      { name: { ref: "Prices!C1" }, stockRole: "high",  categories: "Prices!A2:A4", values: "Prices!C2:C4" },
      { name: { ref: "Prices!D1" }, stockRole: "low",   categories: "Prices!A2:A4", values: "Prices!D2:D4" },
      { name: { ref: "Prices!E1" }, stockRole: "close", categories: "Prices!A2:A4", values: "Prices!E2:E4" }
    ]
  }],
  title: { text: "OHLC" },
  legend: { position: "bottom" },
  styleId: 322
});
return await xlsx.getChart(wb, "Prices", "OHLC");
'''
    return run_xlsx_serve(["exec", str(dst), "--create", "--save", "--code", code])


def witan_multilevel() -> dict[str, Any]:
    dst = OUT / "case24_witan_multilevel.xlsx"
    code = r'''
await xlsx.addSheet(wb, "Data");
await xlsx.setCells(wb, [
  {address:"Data!A1", value:"Region"}, {address:"Data!B1", value:"Quarter"}, {address:"Data!C1", value:"Revenue"},
  {address:"Data!A2", value:"North"},  {address:"Data!B2", value:"Q1"},      {address:"Data!C2", value:120},
  {address:"Data!A3", value:"North"},  {address:"Data!B3", value:"Q2"},      {address:"Data!C3", value:140},
  {address:"Data!A4", value:"South"},  {address:"Data!B4", value:"Q1"},      {address:"Data!C4", value:90},
  {address:"Data!A5", value:"South"},  {address:"Data!B5", value:"Q2"},      {address:"Data!C5", value:110}
]);
await xlsx.addChart(wb, "Data", {
  name: "Revenue by Region Quarter",
  position: { from: { cell: "E2" }, to: { cell: "M18" } },
  groups: [{
    type: "column",
    series: [{
      name: { ref: "Data!C1" },
      categories: "Data!A2:B5",
      categoriesRefType: "multiLevelString",
      values: "Data!C2:C5"
    }]
  }],
  title: { text: "Revenue" },
  legend: { position: "right" },
  axes: { category: { title: { text: "Region / Quarter" } }, value: { title: { text: "Revenue" } } },
  styleId: 201
});
return await xlsx.getChart(wb, "Data", "Revenue by Region Quarter");
'''
    return run_xlsx_serve(["exec", str(dst), "--create", "--save", "--code", code])


def openpyxl_stock() -> Path:
    dst = OUT / "case23_openpyxl_stock.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Prices"
    for row in [
        ["Date", "Open", "High", "Low", "Close"],
        ["2026-01-01", 100, 112, 96, 108],
        ["2026-01-02", 108, 116, 101, 103],
        ["2026-01-03", 103, 120, 99, 117],
    ]:
        ws.append(row)

    chart = StockChart()
    chart.title = "OHLC"
    chart.add_data(Reference(ws, min_col=2, max_col=5, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    ws.add_chart(chart, "G2")
    wb.save(dst)
    return dst


def openpyxl_multilevel() -> Path:
    dst = OUT / "case24_openpyxl_multilevel.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in [
        ["Region", "Quarter", "Revenue"],
        ["North", "Q1", 120],
        ["North", "Q2", 140],
        ["South", "Q1", 90],
        ["South", "Q2", 110],
    ]:
        ws.append(row)

    chart = BarChart()
    chart.type = "col"
    chart.title = "Revenue"
    chart.y_axis.title = "Revenue"
    chart.x_axis.title = "Region / Quarter"
    chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, max_col=2, min_row=2, max_row=5))
    ws.add_chart(chart, "E2")
    wb.save(dst)
    return dst


def waterfall_summary(chart: dict[str, Any] | None) -> dict[str, Any]:
    if chart is None:
        return {"supported": hasattr(openpyxl_charts, "WaterfallChart")}
    series = first_series(chart)
    return {
        "position": pos_range(chart),
        "groupType": first_group(chart).get("type"),
        "totalIndexes": series.get("totalIndexes"),
        "dataLabels": series.get("dataLabels"),
        "valueAxisFormat": chart.get("axes", {}).get("value", {}).get("numberFormat"),
    }


def stock_summary(chart: dict[str, Any]) -> dict[str, Any]:
    group = first_group(chart)
    series = group.get("series") or []
    return {
        "position": pos_range(chart),
        "groupType": group.get("type"),
        "stockRoles": [s.get("stockRole") for s in series],
        "categoryRefTypes": sorted({s.get("categoriesRefType") for s in series}),
        "legendPosition": chart.get("legend", {}).get("position"),
        "styleId": chart.get("styleId"),
    }


def multilevel_summary(chart: dict[str, Any]) -> dict[str, Any]:
    return {
        "position": pos_range(chart),
        "groupType": first_group(chart).get("type"),
        "categories": first_series(chart).get("categories"),
        "categoriesRefType": first_series(chart).get("categoriesRefType"),
        "categoryAxisPosition": chart.get("axes", {}).get("category", {}).get("position"),
    }


def witan_waterfall_expected(summary: dict[str, Any]) -> bool:
    return (
        summary["position"] == "D2:L18"
        and summary["groupType"] == "waterfall"
        and summary["totalIndexes"] == [3, 5]
        and summary["valueAxisFormat"] == "$#,##0"
    )


def witan_stock_expected(summary: dict[str, Any]) -> bool:
    return (
        summary["position"] == "G2:N18"
        and summary["groupType"] == "stockOHLC"
        and summary["stockRoles"] == ["open", "high", "low", "close"]
        and summary["categoryRefTypes"] == ["string"]
        and summary["legendPosition"] == "bottom"
    )


def witan_multilevel_expected(summary: dict[str, Any]) -> bool:
    return (
        summary["position"] == "E2:M18"
        and summary["groupType"] == "column"
        and summary["categoriesRefType"] == "multiLevelString"
    )


def main() -> int:
    shutil.rmtree(OUT, ignore_errors=True)
    OUT.mkdir(parents=True)

    witan_waterfall_result = waterfall_summary(witan_waterfall())
    openpyxl_waterfall_result = waterfall_summary(None)

    witan_stock_result = stock_summary(witan_stock())
    openpyxl_stock_path = openpyxl_stock()
    openpyxl_stock_result = stock_summary(get_chart(openpyxl_stock_path, "Prices", "Chart 1"))
    lo_stock_result = stock_summary(get_chart(resave_with_libreoffice(openpyxl_stock_path), "Prices", "Chart 1"))

    witan_multilevel_result = multilevel_summary(witan_multilevel())
    openpyxl_multilevel_path = openpyxl_multilevel()
    openpyxl_multilevel_result = multilevel_summary(get_chart(openpyxl_multilevel_path, "Data", "Chart 1"))
    lo_multilevel_result = multilevel_summary(
        get_chart(resave_with_libreoffice(openpyxl_multilevel_path), "Data", "Chart 1")
    )

    print("22 Waterfall chart")
    print(f"   witan: {witan_waterfall_result}")
    print(f"   openpyxl: {openpyxl_waterfall_result}")
    print("   openpyxl + LibreOffice: no source chart; openpyxl cannot author waterfall")
    print()

    print("23 Stock OHLC chart")
    print(f"   witan: {witan_stock_result}")
    print(f"   openpyxl: {openpyxl_stock_result}")
    print(f"   openpyxl + LibreOffice: {lo_stock_result}")
    print()

    print("24 Multi-level category labels")
    print(f"   witan: {witan_multilevel_result}")
    print(f"   openpyxl: {openpyxl_multilevel_result}")
    print(f"   openpyxl + LibreOffice: {lo_multilevel_result}")
    print()

    unexpected = (
        not witan_waterfall_expected(witan_waterfall_result)
        or openpyxl_waterfall_result["supported"]
        or not witan_stock_expected(witan_stock_result)
        or witan_stock_expected(openpyxl_stock_result)
        or witan_stock_expected(lo_stock_result)
        or not witan_multilevel_expected(witan_multilevel_result)
        or witan_multilevel_expected(openpyxl_multilevel_result)
        or witan_multilevel_expected(lo_multilevel_result)
    )
    print(f"Summary: {'expected comparisons' if not unexpected else 'unexpected result'}")
    return 1 if unexpected else 0


if __name__ == "__main__":
    raise SystemExit(main())
